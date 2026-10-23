# -*- coding: utf-8 -*-
"""The final presentation, in the Presentation 1 design (le_theme.py), from final.py (NEXT-STEPS.md §7).

No template or rubric exists for the final (KB §50, D2), so the deck follows the final report outline: the question,
the data and the grid, the thermal pair, the landforms, the classification, the verdicts, the limits, what next. It
is built as the interim is (build_interim_live.py), from one content module and from figures the project exports:

  <ArcGIS Pro Python> make_final_figs.py           the layouts, to final_img\\
  py -3.13 make_fig_final.py                       the two charts the interim never drew, to final_img\\
  py -3.13 build_final_deck.py "..\\NEXT STUFF\\Final\\Mars Global Mosaic - Final Presentation.pptx"

The figure-slide placement (a tall figure left with its takeaways right; a wide map full width with its legend beside
the takeaways; tall two-frame sheets map-left) is the interim builder's, copied so that build stays untouched.
"""
import os
import sys

import final as F
from le_theme import (deck, new, title_slide, section, bullets, table, note, textbox, style, no_bullet, bullet,
                      png_size, PT, MARGIN, CONTENT_W, BODY_Y, BODY_H, CYAN, GREY, NAVY, INK)
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "final_img")
OUT = sys.argv[1]
RIGHT = MARGIN + CONTENT_W
FIG_H = 368 * PT
SIDE_MAX_W = 520 * PT
WIDE_MAX_H = 290 * PT
MINUS = "−"


def image_path(kind, ref):
    p = os.path.join(IMG, ref + ".png") if kind == "layout" else os.path.join(HERE, ref)
    if not os.path.exists(p):
        sys.exit("missing image %s: run make_final_figs.py / make_fig_final.py first" % p)
    return p


def add_pic(s, path, x, y, w=None, h=None):
    px_w, px_h = png_size(path)
    if w is None:
        w = int(h * px_w / px_h)
    if h is None:
        h = int(w * px_h / px_w)
    s.shapes.add_picture(path, Emu(x), Emu(y), Emu(w), Emu(h))
    return w, h


def source(s, text, x, y, w):
    tf = textbox(s, x, y, w, 18 * PT)
    p = tf.paragraphs[0]
    no_bullet(p)
    style(p.add_run(), 10, italic=True, color=GREY).text = "PROJECT-KNOWLEDGE.md " + text


def below(s, f, y):
    bullets(s, f["points"], size=14, top=y, height=510 * PT - y, space=4)
    source(s, f["source"], MARGIN, 516 * PT, CONTENT_W)


def split_legend(s, f, path):
    base = path[:-4]
    m, lg = base + "__map.png", base + "__legend.png"
    if not (os.path.exists(m) and os.path.exists(lg)):
        return False
    px_w, px_h = png_size(m)
    if f.get("map_left"):
        w, h = add_pic(s, m, MARGIN, BODY_Y, h=FIG_H)
        x = MARGIN + w + 22 * PT
        tw = RIGHT - x
        bullets(s, f["points"], size=14, top=BODY_Y, height=176 * PT, space=6, left=x, width=tw)
        y = BODY_Y + 184 * PT
        lw_px, lh_px = png_size(lg)
        lh = BODY_Y + FIG_H - 22 * PT - y
        lw = min(int(lh * lw_px / lh_px), tw)
        add_pic(s, lg, x, y, w=lw, h=int(lw * lh_px / lw_px))
        source(s, f["source"], x, BODY_Y + FIG_H - 18 * PT, tw)
        return True
    w = CONTENT_W
    h = int(w * px_h / px_w)
    cap = f.get("map_h", 0) * PT or WIDE_MAX_H
    if h > cap:
        h = cap
        w = int(h * px_w / px_h)
    add_pic(s, m, MARGIN + (CONTENT_W - w) // 2, BODY_Y, w=w, h=h)
    y = BODY_Y + h + 10 * PT
    room = 508 * PT - y
    lw_px, lh_px = png_size(lg)
    lw = min(int(room * lw_px / lh_px), 300 * PT)
    lh = int(lw * lh_px / lw_px)
    add_pic(s, lg, RIGHT - lw, y, w=lw, h=lh)
    bullets(s, f["points"], size=14, top=y, height=room, space=4, width=CONTENT_W - lw - 16 * PT)
    source(s, f["source"], MARGIN, 516 * PT, CONTENT_W)
    return True


def figure_slide(prs, f):
    s = new(prs, f["title"], size=30)
    paths = [image_path(k, r) for k, r in f["images"]]
    px_w, px_h = png_size(paths[0])
    aspect = px_w / float(px_h)
    side_h = f.get("side_h", 0) * PT or FIG_H
    if side_h * aspect <= SIDE_MAX_W or f.get("side_h"):
        w, _ = add_pic(s, paths[0], MARGIN, BODY_Y, h=side_h)
        x = MARGIN + w + 22 * PT
        tw = RIGHT - x
        bullets(s, f["points"], size=15 if tw > 300 * PT else 14, top=BODY_Y + 8 * PT,
                height=320 * PT, space=12, left=x, width=tw)
        source(s, f["source"], x, BODY_Y + side_h - 18 * PT, tw)
    elif split_legend(s, f, paths[0]):
        pass
    else:
        w = CONTENT_W
        h = int(w / aspect)
        if h > WIDE_MAX_H:
            h = WIDE_MAX_H
            w = int(h * aspect)
        add_pic(s, paths[0], MARGIN + (CONTENT_W - w) // 2, BODY_Y, w=w, h=h)
        below(s, f, BODY_Y + h + 10 * PT)
    return s


def labeled(s, items, size=14, top=BODY_Y, height=BODY_H, space=9):
    """le_theme.labeled with ": " between head and text, the project's title form since KB §67.2."""
    tf = textbox(s, MARGIN, top, CONTENT_W, height)
    for i, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        bullet(p)
        style(p.add_run(), size, bold=True, color=NAVY).text = head + ": "
        style(p.add_run(), size, color=INK).text = body
    return tf


def question_slide(prs):
    s = new(prs, "The Question")
    tf = textbox(s, MARGIN, BODY_Y, CONTENT_W, 90 * PT)
    p = tf.paragraphs[0]
    no_bullet(p)
    style(p.add_run(), 18).text = F.QUESTION
    bullets(s, F.GOALS, size=16, top=250 * PT, height=260 * PT, space=10)


def bands_slide(prs):
    s = new(prs, "Relevant Spectral Bands")
    table(s, ["Band or product", "Wavelength", "Cell", "What it contributes"],
          [list(r) for r in F.BANDS], widths=[2.3, 1.9, 0.7, 6.2], size=11, row_h=40 * PT)


def correlation_slide(prs):
    s = new(prs, "Why Fuse Them: Three Independent Dimensions")
    c = F.CORRELATION
    rows = [[b] + [("%+.3f" % v).replace("-", MINUS) if v != 1 else "1" for v in r]
            for b, r in zip(c["bands"], c["r"])]
    t = table(s, ["Ius Chasma, r"] + c["bands"], rows, widths=[1.6] + [1.0] * 5, size=13, row_h=34 * PT)
    for i, r in enumerate(c["r"], 1):          # the independent pairs in cyan, so the eye finds them
        for j, v in enumerate(r, 1):
            if abs(v) < 0.2:
                t.cell(i, j).text_frame.paragraphs[0].runs[0].font.color.rgb = CYAN
    note(s, c["takeaway"] + "   PROJECT-KNOWLEDGE.md §18.3", 360 * PT, size=13)


SPECIAL = {"question": question_slide, "bands": bands_slide, "correlation": correlation_slide}

prs = deck()
counter = iter(range(1, 20))

# ------------------------------------------------------------ title, then the talk in FIGURES order
title_slide(prs, F.DECK_TITLE, (F.TITLE, F.COURSE + " · " + F.AUTHOR,
                                "Final Presentation · " + F.STATUS_DATE))

for f in F.FIGURES:
    if "section" in f:
        head, blurb = f["section"]
        section(prs, "%02d" % next(counter), head, blurb)
    if "special" in f:
        SPECIAL[f["special"]](prs)
    else:
        figure_slide(prs, f)

# ------------------------------------------------------------ verdicts, limits, conclusions
section(prs, "%02d" % next(counter), "What It Adds Up To", "The hypotheses, the limits, the answer")

s = new(prs, "The Hypotheses and Tests")
table(s, ["", "Claim", "Verdict", "KB"], [list(r) for r in F.HYPOTHESES],
      widths=[0.5, 3.6, 6.0, 1.0], size=12, row_h=50 * PT)

s = new(prs, "Limits")
labeled(s, F.LIMITS, size=14, space=8)

s = new(prs, "Conclusions")
bullets(s, F.CONCLUSIONS, size=17, top=BODY_Y, height=BODY_H, space=14)

s = new(prs, "What Next")
bullets(s, F.NEXT_STEPS, size=17, top=BODY_Y, height=260 * PT, space=14)
note(s, F.DATA, 440 * PT, size=13)

s = new(prs, "Questions?", dark=True, size=44)
tf = textbox(s, MARGIN, 220 * PT, CONTENT_W, 120 * PT)
p = tf.paragraphs[0]
no_bullet(p)
style(p.add_run(), 18, color=CYAN).text = F.TITLE
p = tf.add_paragraph()
no_bullet(p)
style(p.add_run(), 14, color=CYAN).text = F.AUTHOR + " · " + F.COURSE

prs.save(OUT)
print("wrote", OUT, "-", len(prs.slides._sldIdLst), "slides")
