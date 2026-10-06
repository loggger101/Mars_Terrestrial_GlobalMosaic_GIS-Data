# -*- coding: utf-8 -*-
"""Slide-building helpers that match the supplied template's typography."""

from lxml import etree
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

# Geometry taken verbatim from the supplied template (EMU).
BANNER = dict(left=3649346, top=428015, width=4792722, height=392159)

MARGIN_L = 838200            # 0.92"
CONTENT_W = 10515600         # 11.5"
TITLE_TOP = 950000
TITLE_H = 520000
RULE_TOP = 1510000
BODY_TOP = 1660000
BODY_H = 4700000

RUST = RGBColor(0xA8, 0x4B, 0x2A)       # accent, keyed to the subject
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x59, 0x59, 0x59)
BAND = RGBColor(0xF2, 0xEE, 0xEB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"


def _sub(parent, tag, **attrs):
    el = etree.SubElement(parent, "{%s}%s" % (A, tag))
    for k, v in attrs.items():
        el.set(k, v)
    return el


def bullet(p, level=0, char="•"):
    """Apply a real DrawingML bullet to a paragraph (python-pptx has no API)."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(228600 + 285750 * level))
    pPr.set("indent", "-228600")
    # child order matters: buFont must precede buChar
    _sub(pPr, "buFont", typeface="Arial", panose="020B0604020202020204",
         pitchFamily="34", charset="0")
    _sub(pPr, "buChar", char=char)


def no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    _sub(pPr, "buNone")


def textbox(slide, left, top, width, height, wrap=True, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def style(run, size=16, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    return run


def add_banner(slide, text):
    tf = textbox(slide, BANNER["left"], BANNER["top"], BANNER["width"],
                 BANNER["height"], wrap=False)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    style(p.add_run(), 18, bold=True)
    p.runs[0].text = text
    return tf


def add_title(slide, text, size=26):
    tf = textbox(slide, MARGIN_L, TITLE_TOP, CONTENT_W, TITLE_H)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    style(r, size, bold=True, color=RUST)
    # thin rule beneath the title
    line = slide.shapes.add_connector(1, Emu(MARGIN_L), Emu(RULE_TOP),
                                      Emu(MARGIN_L + CONTENT_W), Emu(RULE_TOP))
    line.line.color.rgb = RUST
    line.line.width = Pt(1.25)
    return tf


def new_slide(prs, banner, title, title_size=26):
    slide = prs.slides.add_slide(prs.slide_layouts[6])   # Blank
    add_banner(slide, banner)
    add_title(slide, title, title_size)
    return slide


def add_bullets(slide, items, top=BODY_TOP, height=BODY_H, size=16,
                left=MARGIN_L, width=CONTENT_W, space=10):
    """items: list of (level, text) or plain strings (level 0)."""
    tf = textbox(slide, left, top, width, height)
    first = True
    for item in items:
        level, text = item if isinstance(item, tuple) else (0, item)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space)
        p.line_spacing = 1.0
        bullet(p, level, "•" if level == 0 else "–")
        r = p.add_run()
        r.text = text
        style(r, size if level == 0 else size - 2,
              color=INK if level == 0 else GREY)
    return tf


def add_labeled(slide, items, top=BODY_TOP, height=BODY_H, size=15,
                left=MARGIN_L, width=CONTENT_W, space=9):
    """items: list of (bold_label, body_text) rendered as one bulleted run pair."""
    tf = textbox(slide, left, top, width, height)
    first = True
    for label, body in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space)
        p.line_spacing = 1.0
        bullet(p, 0)
        style(p.add_run(), size, bold=True, color=RUST).text = label + " — "
        style(p.add_run(), size, color=INK).text = body
    return tf


def add_table(slide, headers, rows, top=BODY_TOP, height=None, widths=None,
              size=11, header_size=11, left=MARGIN_L, width=CONTENT_W,
              row_h=280000):
    height = height or (row_h * (len(rows) + 1))
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Emu(left),
                                   Emu(top), Emu(width), Emu(height))
    table = shape.table
    table.first_row = True
    table.horz_banding = False

    if widths:
        total = float(sum(widths))
        for i, w in enumerate(widths):
            table.columns[i].width = Emu(int(width * w / total))

    for j, head in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RUST
        cell.margin_left = cell.margin_right = Emu(64008)
        cell.margin_top = cell.margin_bottom = Emu(36576)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        no_bullet(p)
        style(p.add_run(), header_size, bold=True, color=WHITE).text = head

    for i, row in enumerate(rows, start=1):
        table.rows[i].height = Emu(row_h)
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BAND if i % 2 == 0 else WHITE
            cell.margin_left = cell.margin_right = Emu(64008)
            cell.margin_top = cell.margin_bottom = Emu(27432)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            no_bullet(p)
            style(p.add_run(), size, color=INK).text = str(val)
    return table


def add_note(slide, text, top, size=12):
    tf = textbox(slide, MARGIN_L, top, CONTENT_W, 400000)
    p = tf.paragraphs[0]
    no_bullet(p)
    style(p.add_run(), size, italic=True, color=GREY).text = text
    return tf
