# -*- coding: utf-8 -*-
"""Word helpers for the interim report: paragraphs, headings, bullets, shaded tables (KB §45).

Extracted programmatically from build_interim_docx.py (lines from `def para` to the title block,
unchanged) so the living report and the September builder share one typography: Times New Roman
12 pt, 1-inch margins. init(doc) must be called first; the functions write into that document.
"""
from docx.shared import Pt, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SERIF = "Times New Roman"
doc = None


def init(d):
    global doc
    doc = d


def para(text="", bold=False, italic=False, size=12, align=None, space_after=6,
         space_before=0, indent=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if text:
        r = p.add_run(text)
        r.font.name, r.font.size, r.bold, r.italic = SERIF, Pt(size), bold, italic
    return p


def heading(text):
    p = para(text, bold=True, space_before=12, space_after=6)
    p.paragraph_format.keep_with_next = True
    return p


def bullets(items, level=0):
    for it in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(4)
        if isinstance(it, tuple):
            p.add_run(it[0] + ". ").bold = True
            body = it[1]
        else:
            body = it
        p.add_run(body)
        for r in p.runs:
            r.font.name, r.font.size = SERIF, Pt(12)


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def _no_split(row):
    """Stop a row breaking across a page boundary mid-cell."""
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _repeat_header(row):
    """Repeat this row at the top of every continuation page."""
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def table(headers, rows, widths=None, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        shade(cell, "DCDCDC")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.bold, r.font.name, r.font.size = True, SERIF, Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.name, r.font.size = SERIF, Pt(size)
    if widths:
        total, avail = float(sum(widths)), Inches(6.5)
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = int(avail * w / total)
    _repeat_header(t.rows[0])
    for row in t.rows:
        _no_split(row)
    para(space_after=2)
    return t
