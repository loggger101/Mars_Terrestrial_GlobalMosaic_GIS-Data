# -*- coding: utf-8 -*-
"""The living interim presentation (KB §39.3), in the deck design of mars_theme.py (KB §69.11).

Decided 2026-10-08: no fixed slide count, rebuilt as the data moves. Content comes from interim.py only; the
layouts come from interim_img\\, which make_interim_figs.py refreshes from the project. So the build is two
commands:

  "C:\\Program Files\\ArcGIS\\Pro\\bin\\Python\\envs\\arcgispro-py3\\python.exe" make_interim_figs.py
  python build_interim_live.py <out.pptx>

It covers every heading of the course template ('NEXT STUFF\\2 Interim Presentation Template.pptx'): title,
investigators, goals; relevant spectral bands; tasks and percent complete; preliminary results; issues and
hurdles. Until 2026-10-27 it was drawn in le_theme.py, the Presentation 1 design (git holds that version);
the author then asked for the talk on 10 November to be redesigned, the words and numbers unchanged.
"""
import io
import os
import sys

import numpy as np
from PIL import Image

import interim as I
from mars_theme import (deck, new, blank, line, kicker, bullets, table, note, card, dot, shape, textbox, TINT, DEEP,
                        style, no_bullet, hero_image, graticule, veil, credit, footer, pill, PT, W, H, MARGIN,
                        CONTENT_W, BODY_Y, BOTTOM, HEAD, NAVY, INK, MUTED, HAIR, BLUE, SKY, PALE, WHITE)
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "interim_img")
OUT = sys.argv[1]
RIGHT = MARGIN + CONTENT_W                   # right edge of the content area
FIG_H = 380 * PT                             # tallest a figure may be under the title
SIDE_MAX_W = 520 * PT                        # wider than this at FIG_H, and the takeaways go below
WIDE_MAX_H = 290 * PT                        # a full-width map leaves room for four takeaways
LEGEND = PT // 2                             # a legend's size: 0.5 pt per exported pixel on every slide,
                                             # its text about 1.4 x its size in the layout (KB §69.14)
GAP_PX = 24                                  # blank bands inside a layout export close to this (pixels)
MINUS = "−"

# The backgrounds of the dark slides: Ius Chasma from layouts 01 and 02, the map frame only (pixels)
VISIBLE = os.path.join(IMG, "01_visible.png")
NIGHT = os.path.join(IMG, "02_night_ir.png")
CANYON = (300, 20, 1953, 950)                # the canyon and its tributaries
CREDIT_VIS = "Ius Chasma, 271–286°E: Viking MDIM 2.1 colour on this project's 100 m grid"
CREDIT_NIGHT = "Ius Chasma, 271–286°E: THEMIS night IR on this project's 100 m grid"
GRID = graticule(VISIBLE)                    # painted out of both backgrounds; the maps keep theirs


def dark(p):
    """The navy-ground copy from make_interim_dark.py, if there is one (KB §69.13)."""
    d = os.path.join(IMG, "dark", os.path.basename(p))
    return d if os.path.exists(d) else p


def _keep(used, gap):
    """Indices to keep along one axis: every used one; of each blank run inside, `gap`; of the blank
    margins, none."""
    idx = np.flatnonzero(used)
    if not len(idx):
        return np.arange(len(used))
    keep = []
    for i, j in zip(idx[:-1], idx[1:]):
        keep.append(i)
        if j - i - 1 > gap:                  # a blank run longer than gap: its two ends only
            keep.extend(range(i + 1, i + 1 + gap // 2))
            keep.extend(range(j - (gap - gap // 2), j))
        else:
            keep.extend(range(i + 1, j))
    keep.append(idx[-1])
    return np.array(keep)


_TIGHT = {}


def tight(path):
    """A layout export with its blank page margins cut and its blank bands closed to GAP_PX (KB §69.14):
    Pro places the frames on a page, and the page's empty rows (between the ±60° mosaic's two rows of
    panels, above and below the scale bars) took slide room from the maps. In memory, for the deck only.
    Returns PNG bytes and the size."""
    if path not in _TIGHT:
        a = np.asarray(Image.open(path).convert("RGB")).astype(np.int16)
        ink = np.abs(a - a[0, 0]).max(axis=2) > 24
        a = a[_keep(ink.any(axis=1), GAP_PX)][:, _keep(ink.any(axis=0), GAP_PX)]
        buf = io.BytesIO()
        Image.fromarray(a.astype(np.uint8)).save(buf, "PNG")
        _TIGHT[path] = (buf.getvalue(), (a.shape[1], a.shape[0]))
    return _TIGHT[path]


CHIP = (190, 190, 190)                       # the hillshade's grey: a swatch the navy ground would hide sits on it
ITEM_GAP = 40                                # blank columns this wide separate a legend's columns of items (px)


def _ink(a, ground):
    return np.abs(a.astype(np.int16) - ground).max(axis=2) > 24


def _runs(used, gap):
    """(start, stop) of each run of used indices, runs joined across blanks shorter than gap."""
    idx = np.flatnonzero(used)
    if not len(idx):
        return []
    cuts = np.flatnonzero(np.diff(idx) > gap)
    starts = np.r_[idx[0], idx[cuts + 1]]
    stops = np.r_[idx[cuts], idx[-1]] + 1
    return list(zip(starts, stops))


def _contrast(c1, c2):
    def lum(c):
        c = np.asarray(c, float) / 255
        c = np.where(c <= 0.03928, c / 12.92, ((c + 0.055) / 1.055) ** 2.4)
        return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]
    l1, l2 = sorted((lum(c1), lum(c2)), reverse=True)
    return (l1 + 0.05) / (l2 + 0.05)


def _chip(seg, ground):
    """A swatch too dark for the navy ground (a contrast under 2:1) gets the map's grey behind it, so it reads as
    it does on the map; the text beside it is left alone."""
    ink = _ink(seg, ground)
    runs = _runs(ink.any(axis=0), 12)
    if len(runs) < 2:
        return seg
    c0, c1 = runs[0]                         # the swatch: the leftmost run of ink, before the label
    sw = ink[:, c0:c1]
    if _contrast(seg[:, c0:c1][sw].mean(axis=0), ground) >= 2:
        return seg
    r0, r1 = _runs(sw.any(axis=1), 2)[0]
    seg = seg.copy()
    box = (slice(max(r0 - 6, 0), r1 + 6), slice(max(c0 - 6, 0), c1 + 6))
    patch = seg[box]
    patch[~_ink(patch, ground)] = CHIP
    return seg


def restack(path, room_h, max_w):
    """A wide one-row legend (Pro lays the items out in columns across) regrouped into the columns that fit the
    room under the map at the largest scale, reframed; dark swatches on a grey chip (KB §69.16). In memory, like
    tight(); returns the key add_pic and png_size take."""
    a = np.asarray(Image.open(io.BytesIO(tight(path)[0])).convert("RGB"))
    cols, counts = np.unique(a.reshape(-1, 3), axis=0, return_counts=True)
    ground = cols[counts.argmax()].astype(np.int16)
    ink = _ink(a, ground)
    edge = np.flatnonzero(ink.mean(axis=1) > 0.9)       # the frame: rows and columns inked almost the whole way
    frame = a[edge[0], a.shape[1] // 2] if len(edge) else None
    t, b = (edge[edge < 12].max() + 1, edge[edge > a.shape[0] - 12].min()) if len(edge) else (0, a.shape[0])
    edge = np.flatnonzero(ink.mean(axis=0) > 0.9)
    l, r = (edge[edge < 12].max() + 1, edge[edge > a.shape[1] - 12].min()) if len(edge) else (0, a.shape[1])
    inner, iink = a[t:b, l:r], ink[t:b, l:r]
    segs = []
    for c0, c1 in _runs(iink.any(axis=0), ITEM_GAP):
        rr = _runs(iink[:, c0:c1].any(axis=1), 10 ** 6)[0]
        seg = np.empty((rr[1] - rr[0] + 16, c1 - c0 + 16, 3), np.uint8)   # 8 px of ground round each item,
        seg[:] = ground                                                    # so a chip clears the item's edge
        seg[8:-8, 8:-8] = inner[rr[0]:rr[1], c0:c1]
        segs.append(_chip(seg, ground))
    pad, vgap, hgap = 18, 2, ITEM_GAP - 16

    def size(groups):
        ws = [max(s.shape[1] for s in g) for g in groups]
        hs = [sum(s.shape[0] for s in g) + vgap * (len(g) - 1) for g in groups]
        return sum(ws) + hgap * (len(groups) - 1) + 2 * pad, max(hs) + 2 * pad

    best = None
    n = len(segs)
    for mask in range(2 ** max(n - 1, 0)):   # every split of the items, in order, into columns
        groups, g = [], [segs[0]]
        for j in range(1, n):
            if mask >> (j - 1) & 1:
                groups.append(g)
                g = []
            g.append(segs[j])
        groups.append(g)
        w, h = size(groups)
        scale = min(room_h / h, max_w / w, LEGEND / float(PT))
        if best is None or scale > best[0]:
            best = (scale, groups, w, h)
    _, groups, w, h = best
    out = np.empty((h, w, 3), np.uint8)
    out[:] = ground
    x = pad
    for g in groups:
        y = pad
        for sgm in g:
            out[y:y + sgm.shape[0], x:x + sgm.shape[1]] = sgm
            y += sgm.shape[0] + vgap
        x += max(sgm.shape[1] for sgm in g) + hgap
    if frame is not None:
        out[:2], out[-2:], out[:, :2], out[:, -2:] = frame, frame, frame, frame
    buf = io.BytesIO()
    Image.fromarray(out).save(buf, "PNG")
    key = path + "|stacked"
    _TIGHT[key] = (buf.getvalue(), (w, h))
    return key


def is_layout(path):
    """A navy-ground layout export: interim_img\\dark\\<two digits>_<name>[__map|__legend].png."""
    return os.path.basename(os.path.dirname(path)) == "dark" and os.path.basename(path)[:2].isdigit()


def png_size(path):
    if is_layout(path):
        return tight(path)[1]
    with Image.open(path) as im:
        return im.size


def image_path(kind, ref):
    p = os.path.join(IMG, ref + ".png") if kind == "layout" else os.path.join(HERE, ref)
    if not os.path.exists(p):
        sys.exit("missing image %s: run make_interim_figs.py first" % p)
    return dark(p)


def add_pic(s, path, x, y, w=None, h=None):
    """A picture at its own aspect; a layout export goes in tightened (tight())."""
    px_w, px_h = png_size(path)
    if w is None:
        w = int(h * px_w / px_h)
    if h is None:
        h = int(w * px_h / px_w)
    src = io.BytesIO(tight(path)[0]) if is_layout(path) else path
    s.shapes.add_picture(src, Emu(x), Emu(y), Emu(w), Emu(h))
    return w, h


def below(s, f, y):
    """Takeaways under a picture. The KB § in f["source"] is not printed (KB §69.9)."""
    bullets(s, f["points"], size=14, top=y, height=BOTTOM - y, space=4)


def split_legend(s, f, path):
    """If make_interim_figs.py wrote <layout>__map.png and __legend.png, use them: the map
    across the full width, the takeaways under it on the left, the legend on the right."""
    base = os.path.join(IMG, os.path.basename(path)[:-4])
    m, lg = base + "__map.png", base + "__legend.png"
    if not (os.path.exists(m) and os.path.exists(lg)):
        return False
    m, lg = dark(m), dark(lg)
    px_w, px_h = png_size(m)
    if f.get("map_left"):                    # a tall map (16, 17: frame A over frame B): map left at full
        w, h = add_pic(s, m, MARGIN, BODY_Y, h=FIG_H)   # height; takeaways, then the legend, on the right
        x = MARGIN + w + 24 * PT
        tw = RIGHT - x
        ph = f.get("points_h", 176) * PT     # points_h (pt): more room when the takeaways run long
        bullets(s, f["points"], size=14, top=BODY_Y, height=ph, space=6, left=x, width=tw)
        y = BODY_Y + ph + 8 * PT
        lw_px, lh_px = png_size(lg)
        lh = BODY_Y + FIG_H - y
        lw = min(int(lh * lw_px / lh_px), tw, lw_px * LEGEND)
        add_pic(s, lg, x, y, w=lw, h=int(lw * lh_px / lw_px))
        return True
    w = CONTENT_W
    h = int(w * px_h / px_w)
    cap = f.get("map_h", 0) * PT or WIDE_MAX_H   # map_h (pt): a lower cap, when the legend is tall
    if h > cap:
        h = cap
        w = int(h * px_w / px_h)
    if f.get("legend_beside"):               # a squarish map (09): map left, the legend in the room to its right,
        add_pic(s, m, MARGIN, BODY_Y, w=w, h=h)   # the takeaways full width beneath (KB §69.16)
        x = MARGIN + w + 24 * PT
        lw_px, lh_px = png_size(lg)
        lw = min(RIGHT - x, lw_px * LEGEND, int(h * lw_px / lh_px))
        add_pic(s, lg, x, BODY_Y, w=lw, h=int(lw * lh_px / lw_px))
        below(s, f, BODY_Y + h + 12 * PT)
        return True
    add_pic(s, m, MARGIN + (CONTENT_W - w) // 2, BODY_Y, w=w, h=h)
    y = BODY_Y + h + 12 * PT
    room = BOTTOM - y                        # height left above the footer
    lw_px, lh_px = png_size(lg)
    max_w = 300 * PT
    if lw_px > 3 * lh_px:                    # one wide row of items: regrouped to fit the room (KB §69.16)
        lg = restack(lg, room / float(PT), max_w / float(PT))
        lw_px, lh_px = png_size(lg)
    lw = min(int(room * lw_px / lh_px), max_w, lw_px * LEGEND)
    lh = int(lw * lh_px / lw_px)
    add_pic(s, lg, RIGHT - lw, y, w=lw, h=lh)
    bullets(s, f["points"], size=14, top=y, height=room, space=4, width=CONTENT_W - lw - 18 * PT)
    return True


def figure_slide(prs, f):
    s = new(prs, f["title"], f.get("group"))
    if f.get("verdict"):                     # the hypothesis slides: the answer on the kicker's line
        pill(s, f["verdict"])
    paths = [image_path(k, r) for k, r in f["images"]]
    if len(paths) == 1:
        px_w, px_h = png_size(paths[0])
        aspect = px_w / float(px_h)
        if FIG_H * aspect <= SIDE_MAX_W:     # a tall figure: figure left, takeaways right
            w, _ = add_pic(s, paths[0], MARGIN, BODY_Y, h=FIG_H)
            x = MARGIN + w + 24 * PT
            tw = RIGHT - x
            bullets(s, f["points"], size=15 if tw > 300 * PT else 14, top=BODY_Y + 8 * PT,
                    height=320 * PT, space=12, left=x, width=tw)
        elif split_legend(s, f, paths[0]):   # map full width; takeaways and legend beneath
            pass
        else:                                # a wide figure: full width, takeaways beneath
            w = CONTENT_W
            h = int(w / aspect)
            cap = f.get("map_h", 0) * PT or WIDE_MAX_H   # map_h (pt): taller, when the takeaways are few
            if h > cap:
                h = cap
                w = int(h * aspect)
            add_pic(s, paths[0], MARGIN + (CONTENT_W - w) // 2, BODY_Y, w=w, h=h)
            below(s, f, BODY_Y + h + 12 * PT)
    else:                                    # figures side by side, takeaways beneath
        gap = 16 * PT
        w = int((CONTENT_W - gap * (len(paths) - 1)) / len(paths))
        h = 0
        for i, p in enumerate(paths):
            _, h = add_pic(s, p, MARGIN + i * (w + gap), BODY_Y, w=w)
        below(s, f, BODY_Y + h + 12 * PT)
    for sh in s.shapes:                      # alt text: the slide's title, not the image's file name (KB §69.9)
        if sh.shape_type == 13:
            sh._element.nvPicPr.cNvPr.set("descr", f["title"])
    return s


def band_grid(s, cx, y, i, cell=5.5 * PT):
    """The ±60° grid of 15° blocks, 24 x 8, as a tile strip: band i's two rows (north and south) lit, the
    equator a light rule (KB §69.16)."""
    gw, gh = int(24 * cell), int(8 * cell)
    x = int(cx - gw / 2)
    shape(s, MSO_SHAPE.RECTANGLE, x, y, gw, gh, DEEP)
    for r in (3 - i, 4 + i):                 # row 0 is 45-60°N
        shape(s, MSO_SHAPE.RECTANGLE, x, int(y + r * cell), gw, int(cell), BLUE)
    rule = int(0.75 * PT)
    for c in range(1, 24):                   # the gaps between tiles, in the card's colour
        shape(s, MSO_SHAPE.RECTANGLE, int(x + c * cell - rule / 2), y, rule, gh, TINT)
    for r in range(1, 8):
        shape(s, MSO_SHAPE.RECTANGLE, x, int(y + r * cell - rule / 2), gw, rule, MUTED if r == 4 else TINT)
    return gh


def block_scale_slide(prs, b):
    """The size of a 15° block (KB §69.15): one card per band of latitude, the block drawn to scale as its
    trapezoid (equator side down), the nearest state's area as a dashed square of the same area behind it;
    above it, where the band lies on the ±60° grid (KB §69.16)."""
    s = new(prs, b["title"], b["group"])
    gap = 16 * PT
    cw = int((CONTENT_W - 3 * gap) / 4)
    top, ch = BODY_Y, 340 * PT
    k = 0.16 * PT                            # pt per km, one scale for every card
    for i, (band, km2, mi2, w0, w1, state, ratio) in enumerate(I.BLOCK_SCALE):
        x = MARGIN + i * (cw + gap)
        cx = x + cw // 2
        card(s, x, top, cw, ch)
        line(s, x, top + 12 * PT, cw, 18 * PT, band, 14, bold=True, color=SKY, align=PP_ALIGN.CENTER)
        band_grid(s, cx, top + 38 * PT, i)
        yb = top + 234 * PT                  # the equator side of each block
        h = int(I.BLOCK_NS_KM * k)
        ff = s.shapes.build_freeform(int(cx - w0 * k / 2), yb)
        ff.add_line_segments([(int(cx + w0 * k / 2), yb), (int(cx + w1 * k / 2), yb - h),
                              (int(cx - w1 * k / 2), yb - h)], close=True)
        sh = ff.convert_to_shape()
        sh.fill.solid()
        sh.fill.fore_color.rgb = BLUE
        sh.line.fill.background()
        sh.shadow.inherit = False
        side = int(I.STATE_AREA[state][0] ** 0.5 * k)
        sq = shape(s, MSO_SHAPE.RECTANGLE, cx - side // 2, yb - side, side, side)
        sq.line.color.rgb = WHITE
        sq.line.width = Emu(int(1.25 * PT))
        sq.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        line(s, x, yb + 10 * PT, cw, 26 * PT, km2, 20, font=HEAD, color=WHITE, align=PP_ALIGN.CENTER)
        line(s, x, yb + 37 * PT, cw, 16 * PT, mi2, 12, color=MUTED, align=PP_ALIGN.CENTER)
        line(s, x, yb + 58 * PT, cw, 20 * PT, ratio + " × " + state, 16, bold=True, color=SKY,
             align=PP_ALIGN.CENTER)
        line(s, x, yb + 82 * PT, cw, 16 * PT, "dashed: " + state + ", " + I.STATE_AREA[state][1], 11,
             color=MUTED, align=PP_ALIGN.CENTER)
    note(s, b["note"], top + ch + 12 * PT, size=12)
    return s


def mix(a, b, t):
    return RGBColor(*[int(round(a[i] + (b[i] - a[i]) * t)) for i in range(3)])


prs = deck()
no_footer = []                               # slide numbers that carry no footer

# ------------------------------------------------------------ title, investigators
s = blank(prs, dark=True)
hero_image(s, VISIBLE, CANYON, CREDIT_VIS, GRID)
veil(s, [(0, 97), (45, 93), (100, 38)])
kicker(s, "Interim presentation  ·  " + I.COURSE, y=152 * PT, dark=True)
line(s, MARGIN, 170 * PT, 700 * PT, 70 * PT, I.DECK_TITLE, 54, font=HEAD, color=WHITE)
subtitle = I.TITLE[len(I.DECK_TITLE) + 2:] if I.TITLE.startswith(I.DECK_TITLE + ": ") else I.TITLE   # not twice
line(s, MARGIN, 246 * PT, 500 * PT, 56 * PT, subtitle, 20, color=PALE)
line(s, MARGIN, 322 * PT, 500 * PT, 22 * PT, I.AUTHOR, 16, bold=True, color=WHITE)
line(s, MARGIN, 344 * PT, 500 * PT, 18 * PT, "Status as of " + I.STATUS_DATE, 12, color=PALE)
credit(s, CREDIT_VIS)
no_footer.append(len(prs.slides._sldIdLst))

# ------------------------------------------------------------ goals
s = new(prs, "Project Goals", "The question")
line(s, MARGIN, BODY_Y, CONTENT_W, 90 * PT, I.MISSION, 19, font=HEAD, italic=True, color=INK)
gap = 16 * PT
cw = int((CONTENT_W - 3 * gap) / 4)
for i, g in enumerate(I.GOALS):              # the cards closer under the mission, taller, larger text
    x = MARGIN + i * (cw + gap)
    card(s, x, 216 * PT, cw, 236 * PT)
    dot(s, x + 18 * PT, 236 * PT, 32 * PT, i + 1, size=14)
    line(s, x + 18 * PT, 286 * PT, cw - 36 * PT, 160 * PT, g, 16, color=INK)

# ------------------------------------------------------------ relevant spectral bands
s = new(prs, "Relevant Spectral Bands", "The data")
table(s, ["Band or product", "Wavelength", "Cell", "What it contributes"],
      [[c.replace("r −", "r −").replace("| ≤ ", "| ≤ ") for c in r]
       for r in I.BANDS],                                             # "r −0.49" and "|r| ≤ 0.12" unbroken
      widths=[2.4, 2.25, 0.75, 5.6], size=14, row_h=60 * PT, bold_first=True)

s = new(prs, "Why Fuse Them: Three Independent Dimensions", "The data")
c = I.CORRELATION
rows = [[b] + [("%+.3f" % v).replace("-", MINUS) if v != 1 else "1" for v in r]
        for b, r in zip(c["bands"], c["r"])]
t = table(s, ["Ius Chasma, r"] + c["bands"], rows, widths=[1.6] + [1.0] * 5, size=14, row_h=40 * PT,
          bold_first=True)
for i, r in enumerate(c["r"], 1):            # each cell shaded by |r|, one hue: the pale cells are independent
    for j, v in enumerate(r, 1):
        cell, run = t.cell(i, j), t.cell(i, j).text_frame.paragraphs[0].runs[0]
        cell.fill.solid()
        if i == j:
            cell.fill.fore_color.rgb, run.font.color.rgb = HAIR, MUTED
        else:
            k = abs(v) ** 1.2
            cell.fill.fore_color.rgb = mix(TINT, BLUE, k)
            run.font.color.rgb = WHITE if k > 0.55 else (SKY if abs(v) < 0.2 else INK)
            run.font.bold = abs(v) < 0.2
note(s, c["takeaway"], 372 * PT, size=14)
note(s, "Each cell is shaded by the strength of the correlation: the brighter, the more the two bands repeat "
        "each other.", 412 * PT, size=11)

# ------------------------------------------------------------ tasks and percent complete
def cell(s, x, y, w, h, text, bold=False, color=INK, align=PP_ALIGN.LEFT):
    tf = textbox(s, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = align
    no_bullet(p)
    style(p.add_run(), 10, bold=bold, color=color).text = text


# drawn from shapes, not a table: PowerPoint sizes table rows itself, and the bars must sit on their rows
s = new(prs, "Project Tasks and Percent Complete", "Where it stands")
ROW_H = 24 * PT
cols = [32 * PT, 180 * PT, 424 * PT, 152 * PT, 60 * PT]       # 848 pt
xs = [MARGIN + sum(cols[:k]) for k in range(len(cols))]
shape(s, MSO_SHAPE.RECTANGLE, MARGIN, BODY_Y, CONTENT_W, ROW_H, DEEP)
for k, head in enumerate(["#", "Task", "Where it stands", "13 Sep  \u2192  now", "Now"]):
    cell(s, xs[k] + 7 * PT, BODY_Y, cols[k] - 14 * PT, ROW_H, head, bold=True, color=WHITE)
bw = cols[3] - 16 * PT
for i, ((task, now, pct), sep) in enumerate(zip(I.PROGRESS, I.PROGRESS_SEPT), 1):
    y = BODY_Y + i * ROW_H
    if i % 2 == 0:
        shape(s, MSO_SHAPE.RECTANGLE, MARGIN, y, CONTENT_W, ROW_H, TINT)
    for k, text in ((0, str(i)), (1, task), (2, now), (4, "%d%%" % pct if pct is not None else "\u2013")):
        cell(s, xs[k] + 7 * PT, y, cols[k] - 14 * PT, ROW_H, text, bold=(k == 4))
    mid = y + ROW_H // 2
    bx = xs[3] + 8 * PT
    shape(s, MSO_SHAPE.RECTANGLE, bx, mid - 4 * PT, bw, 8 * PT, HAIR)
    if pct:                                  # two shades of one hue: done by 13 Sep, then added since
        shape(s, MSO_SHAPE.RECTANGLE, bx, mid - 4 * PT, int(bw * pct / 100.0), 8 * PT, BLUE)
    if sep:
        shape(s, MSO_SHAPE.RECTANGLE, bx, mid - 4 * PT, int(bw * sep / 100.0), 8 * PT, SKY)
note(s, "Bar: percent complete toward the final. Light blue: done by 13 September; blue: added since; "
        "dark: still to do.",
     BODY_Y + 15 * ROW_H + 10 * PT, size=11)

s = new(prs, "Processing Completed Since the Last Report", "Where it stands")
table(s, ["When", "What", "Outcome"], [list(r[:3]) for r in I.LOG],   # the KB column is not printed (§69.9)
      widths=[1.0, 4.8, 4.5], size=11, row_h=26 * PT, bold_first=True)

# ------------------------------------------------------------ preliminary results
s = blank(prs, dark=True)
hero_image(s, NIGHT, CANYON, CREDIT_NIGHT, GRID)
veil(s, [(0, 94), (60, 86), (100, 60)])
kicker(s, "Preliminary results", y=112 * PT, dark=True)
line(s, MARGIN, 130 * PT, CONTENT_W, 56 * PT, "What the data show so far", 44, font=HEAD, color=WHITE)
line(s, MARGIN, 192 * PT, 620 * PT, 20 * PT,
     "Every map is a layout in the ArcGIS project, exported fresh for this deck.", 14, color=PALE)
tw = int((CONTENT_W - 3 * 24 * PT) / 4)
for i, (num, what) in enumerate(I.HEADLINES):
    x = MARGIN + i * (tw + 24 * PT)
    line(s, x, 300 * PT, tw, 52 * PT, num, 40, font=HEAD, color=SKY)
    line(s, x, 356 * PT, tw - 8 * PT, 60 * PT, what, 13, color=PALE)
credit(s, CREDIT_NIGHT)
no_footer.append(len(prs.slides._sldIdLst))

for f in I.FIGURES:
    if f["title"] == I.BLOCK_SLIDE["before"]:   # the blocks' size, where they are first used
        block_scale_slide(prs, I.BLOCK_SLIDE)
    figure_slide(prs, f)

# ------------------------------------------------------------ issues and hurdles
def issue_cards(s, items, cols, top, h, size):
    gap = 14 * PT
    cw = int((CONTENT_W - (cols - 1) * gap) / cols)
    for k, (label, body) in enumerate(items):
        x = MARGIN + (k % cols) * (cw + gap)
        y = top + (k // cols) * (h + gap)
        card(s, x, y, cw, h)
        tf = textbox(s, x + 16 * PT, y + 12 * PT, cw - 32 * PT, h - 20 * PT)
        p = tf.paragraphs[0]
        no_bullet(p)
        style(p.add_run(), size + 1, bold=True, color=SKY).text = label
        p = tf.add_paragraph()
        no_bullet(p)
        p.space_before = Emu(3 * PT)
        style(p.add_run(), size, color=INK).text = body


n_limits = 6                                 # the measured limits lead (D12); then the work itself
s = new(prs, "Project Issues and Hurdles", "What the measurements limit")
issue_cards(s, I.ISSUES[:n_limits], 2, BODY_Y, int((BOTTOM - BODY_Y - 28 * PT) / 3), 13)
s = new(prs, "Project Issues and Hurdles", "What slows the work")
gap = 14 * PT                                # three items: one full-width band each, label left (KB §69.14)
rh = int((BOTTOM - 12 * PT - BODY_Y - 2 * gap) / 3)
for k, (label, body) in enumerate(I.ISSUES[n_limits:]):
    y = BODY_Y + k * (rh + gap)
    card(s, MARGIN, y, CONTENT_W, rh)
    tf = textbox(s, MARGIN + 28 * PT, y, 250 * PT, rh, anchor=MSO_ANCHOR.MIDDLE)
    no_bullet(tf.paragraphs[0])
    style(tf.paragraphs[0].add_run(), 20, bold=True, color=SKY).text = label
    tf = textbox(s, MARGIN + 300 * PT, y, CONTENT_W - 324 * PT, rh, anchor=MSO_ANCHOR.MIDDLE)
    no_bullet(tf.paragraphs[0])
    style(tf.paragraphs[0].add_run(), 17, color=INK).text = body

# ------------------------------------------------------------ next steps and schedule
s = new(prs, "Next Steps", "For the final")
gap = 16 * PT
cw = int((CONTENT_W - 3 * gap) / 4)
for i, text in enumerate(I.NEXT_STEPS):
    x = MARGIN + i * (cw + gap)
    card(s, x, BODY_Y, cw, 160 * PT)
    dot(s, x + 14 * PT, BODY_Y + 14 * PT, 26 * PT, i + 1, size=12)
    line(s, x + 14 * PT, BODY_Y + 50 * PT, cw - 28 * PT, 124 * PT, text, 13, color=INK)
ty = 346 * PT                               # the schedule as a timeline
shape(s, MSO_SHAPE.RECTANGLE, MARGIN, ty - PT, CONTENT_W, 2 * PT, HAIR)
step = CONTENT_W / float(len(I.SCHEDULE))
for i, (when, what) in enumerate(I.SCHEDULE):
    cx = int(MARGIN + (i + 0.5) * step)
    here = "interim presentation" in what.lower()   # the talk itself
    d = (16 if here else 11) * PT
    shape(s, MSO_SHAPE.OVAL, cx - d // 2, ty - d // 2, d, d, SKY if here else MUTED)
    line(s, int(cx - step / 2 + 6 * PT), ty - 30 * PT, int(step - 12 * PT), 16 * PT, when, 12, bold=True,
         align=PP_ALIGN.CENTER, color=SKY if here else INK)
    line(s, int(cx - step / 2 + 6 * PT), ty + 18 * PT, int(step - 12 * PT), 110 * PT, what, 11,
         align=PP_ALIGN.CENTER, color=MUTED)

footer(prs, I.DECK_TITLE + "  \u00b7  Interim presentation", set(no_footer))

cp = prs.core_properties                      # python-pptx's defaults name its own author (KB §69.9)
cp.author = cp.last_modified_by = I.AUTHOR
cp.title, cp.subject, cp.comments, cp.keywords = I.TITLE, "Interim presentation", "", ""
cp.revision = 1
prs.save(OUT)
print("wrote", OUT, "-", len(prs.slides._sldIdLst), "slides")
