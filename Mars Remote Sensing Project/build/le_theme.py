# -*- coding: utf-8 -*-
"""Logan Edwards' deck design, measured from 'ocean first pres LE.pdf'.

  960 x 540 pt  |  Aptos Display titles, Aptos body
  60 pt title / 18 pt caption / 12 pt sub-line
  white ground, dark panel #0E2841, cyan accent #0E9ED4
"""

import os
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PT = 12700
W, H = 960 * PT, 540 * PT
A = "http://schemas.openxmlformats.org/drawingml/2006/main"

DISPLAY = "Aptos Display"
BODY = "Aptos"

INK = RGBColor(0x00, 0x00, 0x00)
NAVY = RGBColor(0x0E, 0x28, 0x41)
CYAN = RGBColor(0x0E, 0x9E, 0xD4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x55, 0x5F, 0x66)
PALE = RGBColor(0x9E, 0xD8, 0xED)
ROW = RGBColor(0xEE, 0xF3, 0xF7)

MARGIN = 65 * PT
CONTENT_W = 830 * PT
TITLE_Y = 48 * PT
BAR_Y = 104 * PT
BODY_Y = 132 * PT
BODY_H = 380 * PT


def deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    return prs


def _sub(parent, tag, **attrs):
    el = etree.SubElement(parent, "{%s}%s" % (A, tag))
    for k, v in attrs.items():
        el.set(k, v)
    return el


def bullet(p, level=0, char="•"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(200000 + 260000 * level))
    pPr.set("indent", "-200000")
    _sub(pPr, "buFont", typeface="Arial", pitchFamily="34", charset="0")
    _sub(pPr, "buChar", char=char)


def no_bullet(p):
    _sub(p._p.get_or_add_pPr(), "buNone")


def style(run, size, bold=False, color=INK, font=BODY, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return run


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tf = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h)).text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def blank(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY if dark else WHITE
    return s


def accent_bar(s, y=BAR_Y, x=MARGIN, w=74 * PT, h=4 * PT):
    bar = s.shapes.add_shape(1, Emu(x), Emu(y), Emu(w), Emu(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CYAN
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def new(prs, title, dark=False, size=34):
    s = blank(prs, dark)
    tf = textbox(s, MARGIN, TITLE_Y, CONTENT_W, 60 * PT)
    p = tf.paragraphs[0]
    no_bullet(p)
    style(p.add_run(), size, color=WHITE if dark else INK, font=DISPLAY).text = title
    accent_bar(s)
    return s


def title_slide(prs, title, lines):
    """His Presentation 1 title slide, reproduced."""
    s = blank(prs)
    tf = textbox(s, 0, 196 * PT, W, 80 * PT)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    no_bullet(p)
    style(p.add_run(), 60, font=DISPLAY).text = title
    for i, line in enumerate(lines):
        tf = textbox(s, 0, (282 + i * 23) * PT, W, 24 * PT)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        no_bullet(p)
        style(p.add_run(), 12).text = line
    return s


def section(prs, number, title, blurb=None):
    """Navy divider announcing one of the three required sections."""
    s = blank(prs, dark=True)
    tf = textbox(s, MARGIN, 196 * PT, CONTENT_W, 34 * PT)
    p = tf.paragraphs[0]
    no_bullet(p)
    style(p.add_run(), 14, color=CYAN).text = number
    tf = textbox(s, MARGIN, 226 * PT, CONTENT_W, 70 * PT)
    p = tf.paragraphs[0]
    no_bullet(p)
    style(p.add_run(), 48, color=WHITE, font=DISPLAY).text = title
    accent_bar(s, y=300 * PT)
    if blurb:
        tf = textbox(s, MARGIN, 320 * PT, 640 * PT, 60 * PT)
        p = tf.paragraphs[0]
        no_bullet(p)
        style(p.add_run(), 14, color=PALE).text = blurb
    return s


def bullets(s, items, size=16, top=BODY_Y, height=BODY_H, space=11,
            dark=False, left=MARGIN, width=CONTENT_W):
    tf = textbox(s, left, top, width, height)
    first = True
    for it in items:
        lvl, text = it if isinstance(it, tuple) else (0, it)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space)
        bullet(p, lvl, "•" if lvl == 0 else "–")
        style(p.add_run(), size if lvl == 0 else size - 2,
              color=(WHITE if dark else (INK if lvl == 0 else GREY))).text = text
    return tf


def labeled(s, items, size=15, top=BODY_Y, height=BODY_H, space=13, dark=False):
    tf = textbox(s, MARGIN, top, CONTENT_W, height)
    first = True
    for label, body in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space)
        bullet(p)
        style(p.add_run(), size, bold=True,
              color=CYAN if dark else NAVY).text = label + " — "
        style(p.add_run(), size, color=WHITE if dark else INK).text = body
    return tf


def table(s, headers, rows, top=BODY_Y, widths=None, size=11, row_h=26 * PT,
          left=MARGIN, width=CONTENT_W):
    shape = s.shapes.add_table(len(rows) + 1, len(headers), Emu(left), Emu(top),
                               Emu(width), Emu(row_h * (len(rows) + 1)))
    t = shape.table
    t.first_row = True
    t.horz_banding = False
    if widths:
        tot = float(sum(widths))
        for i, w in enumerate(widths):
            t.columns[i].width = Emu(int(width * w / tot))
    for j, h in enumerate(headers):
        c = t.cell(0, j)
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        c.margin_left = c.margin_right = Emu(6 * PT)
        c.margin_top = c.margin_bottom = Emu(3 * PT)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = c.text_frame.paragraphs[0]
        no_bullet(p)
        style(p.add_run(), size, bold=True, color=WHITE).text = h
    for i, row in enumerate(rows, 1):
        t.rows[i].height = Emu(row_h)
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = ROW if i % 2 == 0 else WHITE
            c.margin_left = c.margin_right = Emu(6 * PT)
            c.margin_top = c.margin_bottom = Emu(2 * PT)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            no_bullet(p)
            style(p.add_run(), size, color=INK).text = str(v)
    return t


def caption(s, text, x, y, w, dark=False, size=18, align=PP_ALIGN.LEFT):
    tf = textbox(s, x, y, w, 30 * PT)
    p = tf.paragraphs[0]
    p.alignment = align
    no_bullet(p)
    style(p.add_run(), size, color=WHITE if dark else INK).text = text
    return tf


def note(s, text, y, dark=False, size=11):
    tf = textbox(s, MARGIN, y, CONTENT_W, 40 * PT)
    p = tf.paragraphs[0]
    no_bullet(p)
    style(p.add_run(), size, italic=True,
          color=PALE if dark else GREY).text = text
    return tf


def png_size(path):
    """Pixel size of any image we place -- PNG, JPEG, whatever Pillow reads.

    This used to parse PNG header bytes directly, which silently returned
    nonsense for a JPEG and made picture_box crop it to nothing.
    """
    with Image.open(path) as im:
        return im.size


def picture(s, img_dir, name, x, y, w, h):
    return s.shapes.add_picture(os.path.join(img_dir, name), Emu(x), Emu(y),
                                Emu(w), Emu(h))


def picture_h(s, img_dir, name, x, y, h):
    """Fixed height, width from native aspect -- never stretch a raster."""
    path = os.path.join(img_dir, name)
    px_w, px_h = png_size(path)
    w = int(h * px_w / px_h)
    s.shapes.add_picture(path, Emu(x), Emu(y), Emu(w), Emu(h))
    return w


def picture_box(s, img_dir, name, x, y, w, h, keep="top"):
    """Fill an exact w x h box, cropping the overflow -- never stretch a raster.

    Panels captured at different pane aspects can then sit in one aligned grid.
    `keep` chooses which part of the source survives a vertical crop.
    """
    path = os.path.join(img_dir, name)
    px_w, px_h = png_size(path)
    native, box = px_w / float(px_h), w / float(h)
    pic = s.shapes.add_picture(path, Emu(x), Emu(y), Emu(w), Emu(h))
    if box > native:                      # box is wider: trim height
        excess = 1.0 - native / box
        if keep == "top":
            pic.crop_bottom = excess
        else:
            pic.crop_top = pic.crop_bottom = excess / 2.0
    elif native > box:                    # box is taller: trim width, centred
        excess = 1.0 - box / native
        pic.crop_left = pic.crop_right = excess / 2.0
    return pic
