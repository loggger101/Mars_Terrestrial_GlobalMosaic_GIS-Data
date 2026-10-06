# -*- coding: utf-8 -*-
"""Interim presentation, following '2 Interim Presentation Template.pptx'."""

import sys
from pptx import Presentation
from pptx.util import Emu, Pt

import content as C
from deck_helpers import (new_slide, add_bullets, add_labeled, add_table,
                          add_note, textbox, style, no_bullet, bullet,
                          MARGIN_L, CONTENT_W, BODY_TOP, RUST, INK, GREY, BAND)

TEMPLATE = sys.argv[1]
OUT = sys.argv[2]
BANNER = "Remote Sensing Project – Interim Presentation"

RNS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'

prs = Presentation(TEMPLATE)
slide1, slide2 = prs.slides[0], prs.slides[1]

# ------------------------------------------------- slide 1: fill the skeleton
# The template's three bold labels stay exactly where they are; content goes
# into a right-hand column so the original layout is preserved.
COL_X, COL_W = 4200000, 7000000


def fix_banner(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and "Remote Sensing Project" in sh.text_frame.text:
            sh.text_frame.paragraphs[0].runs[0].text = BANNER


fix_banner(slide1)
fix_banner(slide2)

# Project Title
tf = textbox(slide1, COL_X, 1470858, COL_W, 1000000)
p = tf.paragraphs[0]
no_bullet(p)
style(p.add_run(), 18, bold=True, color=RUST).text = C.TITLE

# Project Investigators
tf = textbox(slide1, COL_X, 2505670, COL_W, 1150000)
first = True
for line in C.INVESTIGATORS:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    no_bullet(p)
    p.space_after = Pt(3)
    style(p.add_run(), 14, color=INK).text = line

# Project Goals
tf = textbox(slide1, COL_X, 3798333, COL_W, 2500000)
first = True
for line in C.GOALS_SHORT:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    bullet(p)
    p.space_after = Pt(8)
    style(p.add_run(), 14, color=INK).text = line

# --------------------------------- slide 2: drop the skeleton, build the deck
rid = prs.slides._sldIdLst[1].get(RNS)
prs.part.drop_rel(rid)
prs.slides._sldIdLst.remove(prs.slides._sldIdLst[1])

# ------------------------------------------------ TASKS AND PERCENT COMPLETE
s = new_slide(prs, BANNER, "Project Tasks and Percent Complete")
overall = round(sum(v for _, v in C.PROGRESS) / len(C.PROGRESS))
rows = [[i, t, "%d%%" % v, "■" * max(1, round(v / 10)) if v else "–"]
        for i, (t, v) in enumerate(C.PROGRESS, 1)]
tbl = add_table(s, ["#", "Task", "Complete", "Progress"], rows,
                top=1600000, widths=[0.5, 6.6, 1.1, 3.3], size=10.5,
                row_h=265000)
# colour the bar column by progress
for i, (_, v) in enumerate(C.PROGRESS, start=1):
    run = tbl.cell(i, 3).text_frame.paragraphs[0].runs[0]
    run.font.color.rgb = RUST if v else GREY
add_note(s, "Overall project completion: approximately %d%%. "
            "Remaining effort is concentrated in classification, digitizing, and "
            "crater counting — the three tasks that depend on the base mosaic "
            "being final." % overall, top=5900000, size=12)

# ------------------------------------------------------------- WORK COMPLETED
s = new_slide(prs, BANNER, "Processing Completed to Date")
add_table(s, ["When", "Tool", "Parameters", "Outcome"],
          [list(r) for r in C.WORKFLOW_LOG],
          top=1600000, widths=[1.2, 2.2, 3.6, 4.5], size=10, row_h=280000)
add_note(s, "Recovered from the geoprocessing lineage and tool logs in the project "
            "geodatabase — this is the actual run history, failures included.",
         top=5500000, size=11)

# ------------------------------------------------------- RELEVANT SPECTRAL BANDS
s = new_slide(prs, BANNER, "Relevant Spectral Bands")
add_table(s, ["Band or product", "Wavelength", "Why this band"],
          [list(r) for r in C.BANDS],
          top=1600000, widths=[2.4, 1.8, 7.3], size=10.5, row_h=300000)

s = new_slide(prs, BANNER, "Relevant Spectral Bands — Rationale by Region")
add_table(s, ["Region", "Wavelength", "What it contributes"],
          [[a, b, c] for a, b, c in C.SPECTRUM],
          top=1620000, widths=[1.2, 1.6, 8.2], size=12, row_h=450000)
add_note(s, "Held today: THEMIS Day IR as the 100 m morphology base, Viking MDIM 2.1 "
            "colour for visible context, and the HRSC/MOLA DEM for slope and drainage. "
            "THEMIS Night IR is the one missing band, and it is what turns the day "
            "mosaic into a thermal inertia measurement.", top=4200000, size=12)

# ----------------------------------------------------------- PRELIMINARY RESULTS
s = new_slide(prs, BANNER, "Preliminary Results (1 of 2) — Data and Measurements")
add_labeled(s, C.PRELIM[:3], top=1650000, size=15, space=16)

s = new_slide(prs, BANNER, "Preliminary Results (2 of 2) — Derived Products")
add_labeled(s, C.PRELIM[3:], top=1650000, size=15, space=16)
add_note(s, C.PRELIM_NOTE, top=5950000, size=11)

# ----------------------------------------------------------- ISSUES AND HURDLES
s = new_slide(prs, BANNER, "Project Issues and Hurdles (1 of 2)")
add_labeled(s, C.ISSUES[:4], top=1650000, size=15, space=16)

s = new_slide(prs, BANNER, "Project Issues and Hurdles (2 of 2)")
add_labeled(s, C.ISSUES[4:], top=1650000, size=15, space=16)

# ------------------------------------------------------------------ NEXT STEPS
s = new_slide(prs, BANNER, "Remaining Work — Immediate Next Steps")
add_bullets(s, C.NEXT_STEPS, top=1650000, height=4700000, size=16, space=12)

s = new_slide(prs, BANNER, "Remaining Schedule")
add_table(s, ["Week", "Dates (2026)", "Activity", "Deliverable"],
          [list(r) for r in C.SCHEDULE[3:]],
          top=1620000, widths=[0.8, 1.8, 6.6, 2.3], size=11, row_h=330000)
add_note(s, "The critical path runs through weeks 5–8: nothing can be classified "
            "until the composite exists, and nothing can be digitized until it is "
            "classified.", top=5850000, size=11)

prs.save(OUT)
print("wrote", OUT, "-", len(prs.slides._sldIdLst), "slides")
