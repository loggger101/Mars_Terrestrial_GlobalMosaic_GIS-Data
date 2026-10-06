# -*- coding: utf-8 -*-
"""Written interim progress report, mirroring the interim presentation template's
section order. Typography matches '1 Project Statement.pdf': Times New Roman 12 pt."""

import sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import content as C

OUT = sys.argv[1]
SERIF = "Times New Roman"

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = SERIF
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.0
for st in ("List Bullet", "List Bullet 2"):
    s = doc.styles[st]
    s.font.name = SERIF
    s.font.size = Pt(12)
    s.paragraph_format.space_after = Pt(4)
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)


def para(text="", bold=False, italic=False, size=12, align=None, space_after=6,
         space_before=0, indent=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if text:
        r = p.add_run(text)
        r.font.name, r.font.size, r.bold, r.italic = SERIF, Pt(size), bold, italic
    return p


def heading(text):
    p = para(text, bold=True, space_before=12, space_after=6)
    p.paragraph_format.keep_with_next = True
    return p


def bullets(items, level=0):
    for it in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(4)
        if isinstance(it, tuple):
            p.add_run(it[0] + ". ").bold = True
            body = it[1]
        else:
            body = it
        p.add_run(body)
        for r in p.runs:
            r.font.name, r.font.size = SERIF, Pt(12)


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def _no_split(row):
    """Stop a row breaking across a page boundary mid-cell."""
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _repeat_header(row):
    """Repeat this row at the top of every continuation page."""
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def table(headers, rows, widths=None, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        shade(cell, "DCDCDC")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.bold, r.font.name, r.font.size = True, SERIF, Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.name, r.font.size = SERIF, Pt(size)
    if widths:
        total, avail = float(sum(widths)), Inches(6.5)
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = int(avail * w / total)
    _repeat_header(t.rows[0])
    for row in t.rows:
        _no_split(row)
    para(space_after=2)
    return t


overall = round(sum(v for _, v in C.PROGRESS) / len(C.PROGRESS))

# ============================================================== TITLE BLOCK ==
para("Remote Sensing Project – Interim Progress Report", bold=True, size=14,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para(C.TITLE, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
for line in C.INVESTIGATORS:
    para(line, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para(space_after=6)

# ==================================================================== GOALS ==
heading("Project goals")
para(C.MISSION)
bullets(C.GOALS_SHORT)

# ============================================== TASKS AND PERCENT COMPLETE ==
heading("List of project tasks and percent complete")
para("Overall completion stands at approximately %d%%. The base mosaic and its "
     "derived rasters are the gating tasks: classification, digitizing, and crater "
     "counting all wait on them, which is why the lower half of the table is still "
     "in single and low double digits." % overall)
table(["#", "Task", "Percent complete"],
      [[i, t, "%d%%" % v] for i, (t, v) in enumerate(C.PROGRESS, 1)],
      widths=[0.4, 5.4, 1.2], size=11)

# ============================================================ WORK COMPLETED ==
heading("Processing completed to date")
para("Recovered from the geoprocessing lineage and tool logs held in the project "
     "geodatabase. This is the actual run history, failures included.")
table(["When", "Tool", "Parameters", "Outcome"],
      [list(r) for r in C.WORKFLOW_LOG], widths=[0.8, 1.3, 2.0, 2.4], size=9)

# ========================================================= SPECTRAL BANDS ==
heading("Relevant spectral bands")
para("The band set is unchanged from the prospectus. Each of the three landform "
     "families is addressed by at least two independent lines of evidence — one "
     "morphological and one spectral or thermophysical — so that no origin is "
     "assigned on image texture alone.")
table(["Band or product", "Wavelength", "Why this band"],
      [list(r) for r in C.BANDS], widths=[1.6, 1.1, 4.3], size=9.5)
_lbl = para("By spectral region:", bold=True, space_before=6, space_after=4)
_lbl.paragraph_format.keep_with_next = True
table(["Region", "Wavelength", "What it contributes"],
      [[a, b, c] for a, b, c in C.SPECTRUM], widths=[0.8, 1.2, 4.5], size=10)

# ===================================================== PRELIMINARY RESULTS ==
heading("Preliminary results")
para(C.PRELIM_NOTE, italic=True, size=11)
bullets(C.PRELIM)

# ====================================================== ISSUES AND HURDLES ==
heading("Project issues and hurdles")
bullets(C.ISSUES)

# ============================================================= NEXT STEPS ==
heading("Remaining work and schedule status")
bullets(C.NEXT_STEPS)
para("Remaining schedule:", bold=True, space_before=8, space_after=4)
table(["Week", "Dates (2026)", "Activity", "Deliverable"],
      [list(r) for r in C.SCHEDULE[3:]], widths=[0.5, 1.1, 3.6, 1.3], size=10)
para("The critical path runs through weeks 9 to 12. Crater counting cannot begin "
     "until the digitized unit boundaries are final, so any slip in classification "
     "propagates directly into the chronology and, from there, into the final report.",
     italic=True, size=11)

doc.save(OUT)
print("wrote", OUT)
