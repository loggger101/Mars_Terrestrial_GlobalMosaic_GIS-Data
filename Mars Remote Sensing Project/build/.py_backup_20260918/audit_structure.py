# -*- coding: utf-8 -*-
"""Structural audit of the deck.

  1. Does it still cover all eleven bullets of the prospectus template, in
     order, after two slides were added and the HiRISE material removed?
  2. Are the speaker notes still consistent with the slides they sit behind?
     I wrote them before the rework, so they can carry stale claims.
  3. What fonts does the deck actually ask for, and are they installed?
"""
import glob
import os
import re
import zipfile

from pptx import Presentation

DECK = (r"C:\Users\Loggg\Downloads\Mars Remote Sensing Project"
        r"\Mars Global Mosaic - Presentation 1.pptx")

TEMPLATE = [
    ("Introduction", None),
    ("Background: statement of goals or mission", "Background"),
    ("Tasks required", "Tasks Required"),
    ("Significance / importance", "Significance"),
    ("Methods", None),
    ("Which part of the EM spectrum", "Electromagnetic Spectrum"),
    ("Potential remote sensing data sources", "Data Sources"),
    ("Which spectral bands?", "Spectral Bands"),
    ("Resolution required", "Resolution Required"),
    ("Project schedule", "Project Schedule"),
    ("Expected results", None),
    ("Hypotheses to be tested", "Hypotheses"),
    ("Questions to be answered", "Questions"),
    ("Potential problems", "Potential Problems"),
]

prs = Presentation(DECK)
titles = []
for i, s in enumerate(prs.slides, 1):
    txts = [sh.text_frame.text.strip() for sh in s.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()]
    titles.append((i, txts[0] if txts else ""))

print("=" * 80)
print("TEMPLATE COVERAGE (1 Project Statement.pdf)")
joined = " | ".join(t for _i, t in titles)
last_pos = -1
ordered = True
for label, needle in TEMPLATE:
    if needle is None:
        print("   %-42s (section)" % label)
        continue
    hits = [i for i, t in titles if needle.lower() in t.lower()]
    pos = hits[0] if hits else -1
    if pos != -1 and pos < last_pos:
        ordered = False
    if pos != -1:
        last_pos = pos
    print("   %-42s %s" % (label,
                           "slides %s" % hits if hits else "*** MISSING ***"))
print("\n   order preserved: %s" % ordered)

print()
print("=" * 80)
print("SPEAKER NOTES vs SLIDE CONTENT")
STALE = [
    ("HiRISE", "no HiRISE imagery is held"),
    ("re-acquire", "nothing is being re-acquired"),
    ("ground truth", "there is no ground truth"),
    ("validation site", "Jezero is not a validation site"),
    ("type areas", "there is one type area and one context map"),
    ("four attempts", "there were three composite attempts"),
    ("pyramids built", "pyramids are not built"),
    ("32 GB", "the three rasters are 44 GB"),
]
found = 0
for i, s in enumerate(prs.slides, 1):
    note = s.notes_slide.notes_text_frame.text
    for needle, why in STALE:
        if needle.lower() in note.lower():
            found += 1
            print("   slide %-3d %-16s %s" % (i, repr(needle), why))
if not found:
    print("   no stale claims in any of the %d notes" % len(titles))

# every slide should have a note, and it should not simply repeat the title
print()
for i, s in enumerate(prs.slides, 1):
    note = s.notes_slide.notes_text_frame.text.strip()
    title = titles[i - 1][1]
    if not note:
        print("   slide %d has no note" % i)
    elif note.strip().lower() == title.strip().lower():
        print("   slide %d note just repeats the title" % i)

print()
print("=" * 80)
print("FONTS REQUESTED BY THE DECK")
asked = set()
with zipfile.ZipFile(DECK) as z:
    for n in z.namelist():
        if n.startswith("ppt/slides/slide") and n.endswith(".xml"):
            xml = z.read(n).decode("utf-8")
            asked |= set(re.findall(r'typeface="([^"]+)"', xml))
asked = {a for a in asked if not a.startswith("+")}

fonts_dirs = [r"C:\Windows\Fonts",
              os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\Windows\Fonts")]
installed = []
for d in fonts_dirs:
    installed += [os.path.basename(p).lower()
                  for p in glob.glob(os.path.join(d, "*.tt*"))]
blob = " ".join(installed)
for f in sorted(asked):
    key = f.lower().replace(" ", "")
    print("   %-18s %s" % (f, "installed" if key in blob.replace(" ", "")
                           else "NOT FOUND -- will substitute"))
