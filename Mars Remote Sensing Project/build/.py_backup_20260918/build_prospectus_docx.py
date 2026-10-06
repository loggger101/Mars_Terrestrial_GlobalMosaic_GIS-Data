# -*- coding: utf-8 -*-
"""Written project prospectus. Typography matches '1 Project Statement.pdf':
Times New Roman 12 pt, centred title, bulleted lists under plain section labels."""

import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import content as C

OUT = sys.argv[1]
SERIF = "Times New Roman"
INK = RGBColor(0, 0, 0)

doc = Document()

# --- global typography: Times New Roman 12 pt, 1" margins (matches source) ---
normal = doc.styles["Normal"]
normal.font.name = SERIF
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.0

for st in ("List Bullet", "List Bullet 2"):
    s = doc.styles[st]
    s.font.name = SERIF
    s.font.size = Pt(12)
    s.paragraph_format.space_after = Pt(4)

for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)


def para(text="", bold=False, italic=False, size=12, align=None, space_after=6,
         space_before=0, indent=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if text:
        r = p.add_run(text)
        r.font.name = SERIF
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
    return p


def heading(text):
    """Section label: plain-weight in the source, bolded here for navigability."""
    p = para(text, bold=True, space_before=12, space_after=6)
    p.paragraph_format.keep_with_next = True
    return p


def sub(text):
    p = para(text, bold=True, italic=True, space_before=10, space_after=4)
    p.paragraph_format.keep_with_next = True
    return p


def bullets(items, level=0):
    for it in items:
        style = "List Bullet" if level == 0 else "List Bullet 2"
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(4)
        if isinstance(it, tuple):
            label, body = it
            r = p.add_run(label + ". ")
            r.bold = True
        else:
            body = it
        r = p.add_run(body)
        for r in p.runs:
            r.font.name = SERIF
            r.font.size = Pt(12)


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def _no_split(row):
    """Stop a row breaking across a page boundary mid-cell."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _repeat_header(row):
    """Repeat this row at the top of every continuation page."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))


def table(headers, rows, widths=None, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        shade(cell, "DCDCDC")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.bold = True
        r.font.name = SERIF
        r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.name = SERIF
            r.font.size = Pt(size)
    if widths:
        total = float(sum(widths))
        avail = Inches(6.5)
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = int(avail * w / total)
    _repeat_header(t.rows[0])
    for row in t.rows:
        _no_split(row)
    para(space_after=2)
    return t


# ============================================================== TITLE BLOCK ==
para("Remote Sensing Project Prospectus", bold=True, size=14,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para(C.TITLE, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
for line in C.INVESTIGATORS:
    para(line, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para(space_after=6)

# =============================================================== INTRODUCTION
heading("Introduction:")

sub("Background: statement of goals and mission")
para(C.MISSION)
bullets(C.BACKGROUND)

sub("Tasks required to accomplish the goal and meet the mission")
para("Ten tasks carry the project from the data already on disk to a finished landform "
     "map. They are ordered by dependency, not by effort. Tasks 1 and 2 are complete. "
     "Task 3 — putting every layer into one coordinate system — is the gate: "
     "the band composite in task 5 has already failed three times without it, and "
     "everything downstream of task 5 waits on that composite.")
table(["#", "Task", "What it involves"],
      [[i, t, d] for i, (t, d) in enumerate(C.TASKS, 1)],
      widths=[0.4, 2.0, 5.0])

sub("Significance and importance of the mission")
bullets(C.SIGNIFICANCE)

# ==================================================================== METHODS
doc.add_page_break()
heading("Methods")

sub("Which part of the electromagnetic spectrum can be applied in the mission")
para("Four regions of the spectrum carry usable information about the three target "
     "landform families. Each answers a different question, and the project depends "
     "on using them together rather than choosing among them.")
table(["Region", "Wavelength", "What it contributes to the mission"],
      [[a, b, c] for a, b, c in C.SPECTRUM], widths=[0.8, 1.2, 4.5])
bullets(C.SPECTRUM_NOTES)

sub("Potential remote sensing data sources")
para("Every dataset below is public and free, served as GeoTIFF, which ArcGIS Pro "
     "ingests directly — no ISIS preprocessing is required. The first table is "
     "what is already downloaded and loaded into the project; the properties given "
     "were measured from the files themselves rather than quoted from the product "
     "documentation.")
table(["Dataset", "Mission", "Type", "Resolution", "Measured properties"],
      [list(r) for r in C.DATA_HELD], widths=[1.6, 1.0, 1.2, 0.8, 2.4], size=9)
para("The following are not yet acquired, and the project plan depends on them:",
     space_before=6)
table(["Dataset", "Resolution", "Why it is needed"],
      [list(r) for r in C.DATA_NEEDED], widths=[1.5, 0.8, 4.2], size=9.5)

sub("Which spectral bands")
para("The band set is chosen so that each of the three landform families is "
     "addressed by at least two independent lines of evidence — one morphological "
     "and one spectral or thermophysical.")
table(["Band or product", "Wavelength", "Why this band"],
      [list(r) for r in C.BANDS], widths=[1.6, 1.1, 4.3], size=9.5)

sub("Resolution required — spatial, spectral, radiometric, temporal")
for kind, lines in C.RESOLUTION:
    lbl = para(kind, bold=True, space_before=6, space_after=2, indent=0.25)
    lbl.paragraph_format.keep_with_next = True      # never strand the label
    bullets(lines, level=1)

sched = sub("Project schedule")
sched.paragraph_format.page_break_before = True     # 12-row table, own page
table(["Week", "Dates (2026)", "Activity", "Deliverable"],
      [list(r) for r in C.SCHEDULE], widths=[0.5, 1.45, 3.25, 1.3], size=10)
para("Semester dates are assumed and should be adjusted to the syllabus. Weeks 1 "
     "to 3 are complete. The critical path runs through weeks 4 to 8: projecting "
     "every layer into one coordinate system unblocks the band composite, and the "
     "composite unblocks classification, digitizing, and everything downstream of "
     "them.", italic=True, size=11)

# =========================================================== EXPECTED RESULTS
doc.add_page_break()
heading("Expected results")

sub("Hypotheses to be tested")
bullets(C.HYPOTHESES)

sub("Questions to be answered")
bullets(C.QUESTIONS)

sub("Potential problems that may be encountered")
para("Each problem below is paired with the mitigation already planned for it. "
     "Several are expected to constrain the final result rather than be eliminated; "
     "where that is true, it is stated.")
bullets(C.PROBLEMS)

doc.save(OUT)
print("wrote", OUT)
