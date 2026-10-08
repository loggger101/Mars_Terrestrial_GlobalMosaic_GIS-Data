# -*- coding: utf-8 -*-
"""The living interim report: the same content and figures as the living deck (KB §45).

Reads interim.py (every word and number, each with its KB §) and the images make_interim_figs.py
exports, so the report and the deck cannot disagree. Section order follows the course template's
headings, as build_interim_docx.py did; typography from docx_kit.py (Times New Roman 12 pt).

  python build_interim_report_live.py <out.docx>

build_interim_docx.py, which wrote the 18 Sep report from content.py, is left as it was.
"""
import os
import sys
import time

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import interim as I
import docx_kit as K

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "interim_img")
OUT = sys.argv[1]

doc = Document()
K.init(doc)
normal = doc.styles["Normal"]
normal.font.name = K.SERIF
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), K.SERIF)
normal.paragraph_format.space_after = Pt(6)
for st in ("List Bullet", "List Bullet 2"):
    s = doc.styles[st]
    s.font.name = K.SERIF
    s.font.size = Pt(12)
    s.paragraph_format.space_after = Pt(4)
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)


def image_path(kind, ref):
    p = os.path.join(IMG, ref + ".png") if kind == "layout" else os.path.join(HERE, ref)
    if not os.path.exists(p):
        sys.exit("missing image %s: run make_interim_figs.py first" % p)
    return p


def source(text):
    K.para("Source: PROJECT-KNOWLEDGE.md " + text, italic=True, size=10, space_after=10)


def figure(f, n):
    K.heading("Figure %d. %s" % (n, f["title"]))
    paths = [image_path(k, r) for k, r in f["images"]]
    if len(paths) == 1:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(paths[0], width=Inches(6.5))
    else:
        t = doc.add_table(rows=1, cols=len(paths))
        w = Inches(6.5 / len(paths) - 0.05)
        for c, path in zip(t.rows[0].cells, paths):
            c.paragraphs[0].add_run().add_picture(path, width=w)
    K.bullets(f["points"])
    source(f["source"])


# ------------------------------------------------------------ title block
K.para(I.TITLE, bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
K.para("%s  ·  %s" % (I.COURSE, I.AUTHOR), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
K.para("Interim Report  ·  status as of %s" % I.STATUS_DATE, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
       space_after=14)

# ------------------------------------------------------------ investigators, goals
K.heading("1. Project Investigators")
K.para("%s, %s." % (I.AUTHOR, I.COURSE))
K.heading("2. Project Goals")
K.para(I.MISSION)
K.bullets(I.GOALS)

# ------------------------------------------------------------ spectral bands
K.heading("3. Relevant Spectral Bands")
K.table(["Band or product", "Wavelength", "Cell", "What it contributes"], [list(r) for r in I.BANDS],
        widths=[2.2, 1.6, 0.6, 4.2], size=10)
c = I.CORRELATION
K.para("Correlation between the five image bands over the Ius Chasma type area (PROJECT-KNOWLEDGE.md §18.3):",
       space_after=4)
K.table(["r"] + c["bands"], [[b] + [("%+.3f" % v).replace("-", "−") if v != 1 else "1" for v in r]
                             for b, r in zip(c["bands"], c["r"])], widths=[1.4] + [1.0] * 5, size=10)
K.para(c["takeaway"])

# ------------------------------------------------------------ tasks
K.heading("4. Project Tasks and Percent Complete")
K.table(["#", "Task", "Where it stands", "13 Sep", "Now"],
        [[i, t, now, "%d%%" % sep, "%d%%" % pct if pct is not None else "–"]
         for i, ((t, now, pct), sep) in enumerate(zip(I.PROGRESS, I.PROGRESS_SEPT), 1)],
        widths=[0.3, 2.0, 3.6, 0.6, 0.6], size=10)
K.heading("5. Processing Completed Since the Last Report")
K.table(["When", "What", "Outcome", "KB"], [list(r) for r in I.LOG], widths=[0.9, 2.8, 2.4, 0.7], size=10)

# ------------------------------------------------------------ preliminary results
K.heading("6. Preliminary Results")
K.para("Every map is a layout in the ArcGIS project, exported for this report; every number is quoted from "
       "the project knowledge base, whose section is given under each figure.")
for n, f in enumerate(I.FIGURES, 1):
    figure(f, n)

# ------------------------------------------------------------ issues, next steps
K.heading("7. Project Issues and Hurdles")
K.bullets([(a, b) for a, b in I.ISSUES])
K.heading("8. Next Steps")
K.bullets(I.NEXT_STEPS)
K.table(["When", "What"], [list(r) for r in I.SCHEDULE], widths=[1.4, 5.1], size=10)
K.para("Generated %s by build_interim_report_live.py from interim.py, the same source as the interim "
       "presentation." % time.strftime("%Y-%m-%d"), italic=True, size=10)

doc.save(OUT)
print("wrote", OUT, "-", len(I.FIGURES), "figures")
